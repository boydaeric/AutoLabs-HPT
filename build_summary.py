#!/usr/bin/env python3
"""Build autonomic_rate_summary.xlsx from hpt_autonomic_rates.csv.

Every statistic in the workbook is an Excel formula over the Rates_data tab
(or a link to a Reference input), so the book recalculates if inputs change.
Run scripts/recalc.py (xlsx skill) afterwards to populate cached values.

Usage:  python build_summary.py [--rates hpt_autonomic_rates.csv]
            [--addb reference/<Addendum B>.xlsx] [--out autonomic_rate_summary.xlsx]
"""
import argparse
import re

import openpyxl
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

CODES = ["95921", "95922", "95923", "95924", "93660"]
BATTERY = ("95924", "95923")
HOSPITALS = [  # (name in data, short, clinician label for reconciliation)
    ("Massachusetts General Hospital", "MGH", "MGH (Robbins)"),
    ("Brigham and Women's Hospital", "BWH", None),
    ("Brigham and Women's Faulkner Hospital", "BWFH", "Faulkner (Novak)"),
    ("Wentworth-Douglass Hospital", "WDH", None),
    ("Beth Israel Deaconess Medical Center", "BIDMC", "BIDMC (Freeman + Gibbons)"),
]
SHORT = {h: s for h, s, _ in HOSPITALS}

# CY2026 MPFS, Metro Boston locality 1421201 -- values supplied by the user.
MPFS = {"95921": (54.93, 47.84, 102.77), "95922": (54.14, 49.32, 103.46),
        "95923": (94.02, 47.34, 141.36), "95924": (79.56, 95.23, 174.79)}
MPFS_BATTERY_GIVEN = {"global": 316.15, "technical": 173.58, "professional": 142.57}
# CY2024 Medicare claims patient counts -- supplied by the user.
PATIENTS_2024 = {("MGH", "95924"): 94, ("MGH", "95923"): 95,
                 ("BWFH", "95924"): 114, ("BWFH", "95923"): 114,
                 ("BIDMC", "95924"): 99, ("BIDMC", "95923"): 99}
ADDB_URL = ("https://www.cms.gov/medicare/payment/prospective-payment-systems/"
            "hospital-outpatient-pps/quarterly-addenda-updates/july-2026-addendum-b")

FOCUS = [  # (label, criteria column, criteria value, segment filter)
    ("Blue Cross Blue Shield of Massachusetts", "F", "Blue Cross Blue Shield of MA", "Commercial"),
    ("Harvard Pilgrim / Point32Health", "F", "Harvard Pilgrim / Point32Health", "Commercial"),
    ("Tufts Health Plan", "F", "Tufts Health Plan", "Commercial"),
    ("Mass General Brigham Health Plan", "F", "Mass General Brigham Health Plan", "Commercial"),
    ("UnitedHealthcare", "F", "UnitedHealthcare", "Commercial"),
    ("Aetna", "F", "Aetna", "Commercial"),
    ("Cigna", "F", "Cigna", "Commercial"),
    ("MassHealth managed-care plans", "G", "MassHealth / dual managed care", None),
    ("Medicare Advantage plans", "G", "Medicare Advantage", None),
]

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=13, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
FILL_HEAD = PatternFill("solid", fgColor="305496")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_FLAG = PatternFill("solid", fgColor="FCE4D6")
USD = '$#,##0.00;($#,##0.00);"-"'
MULT = '0.00"x"'
PCT = '0.0%;(0.0%);"-"'


# --------------------------------------------------------------------------
# Classification
# --------------------------------------------------------------------------
def payer_family(name):
    n = name.upper()
    rules = [("BLUE CROSS", "Blue Cross Blue Shield of MA"), ("BCBS", "Blue Cross Blue Shield of MA"),
             ("HARVARD PILGRIM", "Harvard Pilgrim / Point32Health"),
             ("TUFTS HEALTH PUBLIC", "Tufts Health Public Plans"),
             ("TUFTS", "Tufts Health Plan"), ("MGB HEALTH PLAN", "Mass General Brigham Health Plan"),
             ("UNITED", "UnitedHealthcare"), ("AETNA", "Aetna"), ("CIGNA", "Cigna"),
             ("COMMONWEALTH CARE ALLIANCE", "Commonwealth Care Alliance"),
             ("MASSHEALTH", "MassHealth"), ("WELLSENSE", "WellSense"),
             ("MEDICARE", "Medicare"), ("TRICARE", "TRICARE"), ("UNICARE", "Unicare (GIC)"),
             ("WELLPOINT", "Wellpoint"), ("HEALTH NEW ENGLAND", "Health New England"),
             ("HUMANA", "Humana"), ("MARTIN", "Martin's Point"),
             ("CENTERS OF EXCELLENCE", "Centers of Excellence"),
             ("INTERNATIONAL", "International commercial")]
    for key, fam in rules:
        if key in n:
            return fam
    return re.sub(r"\s*\[\d+\]$", "", name).title()


def segment(payer, plan):
    p, pl = payer.upper(), plan.upper()
    if re.match(r"MEDICARE \[2001\]", p):
        return "Medicare FFS"
    if "TRICARE" in p:
        return "TRICARE"
    if re.search(r"MEDICARE|MCARE", pl):
        return "Medicare Advantage"
    if "COMMONWEALTH CARE ALLIANCE" in p or "MASSHEALTH" in p or "WELLSENSE" in p:
        return "MassHealth / dual managed care"
    return "Commercial"


def exclusion(r):
    if not r.payer_name and not r.plan_name:
        return "no payer: gross/discounted-cash only row (not a negotiated rate)"
    if r.setting == "inpatient":
        return "inpatient setting"
    if r.billing_class not in ("", "facility", "both"):
        return f"billing class {r.billing_class}"
    if r.implied_dollar == "":
        return "no dollar or percentage-derived rate"
    return ""


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def style_header(ws, row, ncols, height=30):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = F_HEAD, FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def put(ws, row, col, value, font=F_BASE, fmt=None, fill=None, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def title(ws, text, sub=None):
    put(ws, 1, 1, text, F_TITLE)
    if sub:
        put(ws, 2, 1, sub, F_NOTE)


def fnum(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load_addb(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    out = {}
    for row in wb.worksheets[0].iter_rows(values_only=True):
        if row and str(row[0]).strip() in CODES:
            out[str(row[0]).strip()] = dict(desc=row[1], si=row[2], apc=str(row[3]), rate=row[5])
    missing = set(CODES) - set(out)
    if missing:
        raise SystemExit(f"Addendum B missing codes: {missing}")
    return out


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rates", default="hpt_autonomic_rates.csv")
    ap.add_argument("--addb", default="reference/2026 July Web Addendum B.07.14.26.xlsx")
    ap.add_argument("--out", default="autonomic_rate_summary.xlsx")
    args = ap.parse_args()

    df = pd.read_csv(args.rates, dtype=str, keep_default_na=False)
    df["code"] = df.code.str.strip()
    df["family"] = df.payer_name.map(payer_family)
    df["segment"] = [segment(p, pl) if (p or pl) else "" for p, pl in zip(df.payer_name, df.plan_name)]
    df["exclusion"] = [exclusion(r) for r in df.itertuples()]
    df["included"] = (df.exclusion == "").astype(int)
    df["pct_based"] = ((df.methodology == "percent of total billed charges")
                       | (df.standard_charge_percentage != "")).astype(int)
    df = df.sort_values(["hospital", "code", "segment", "payer_name", "plan_name"],
                        key=lambda s: s.map({h: i for i, (h, _, _) in enumerate(HOSPITALS)})
                        if s.name == "hospital" else s).reset_index(drop=True)
    addb = load_addb(args.addb)

    wb = openpyxl.Workbook()
    ws_notes = wb.active
    ws_notes.title = "Notes"
    names = ["By_code_payer", "Battery", "Multiples", "Focus_payers", "Coverage",
             "Reconciliation", "Medicare_check", "QA", "Reference", "Rates_data"]
    S = {n: wb.create_sheet(n) for n in names}

    # ---------------------------------------------------------------- Reference
    ws = S["Reference"]
    title(ws, "Reference prices",
          "Blue = input values. MPFS values supplied by the user; OPPS values from CMS Addendum B.")
    put(ws, 4, 1, "CY2026 Medicare Physician Fee Schedule — Metro Boston locality 1421201 "
                  "(values supplied by user)", F_BOLD)
    hdr = ["Code", "Technical (TC)", "Professional (PC)", "Global"]
    for i, h in enumerate(hdr, 1):
        put(ws, 5, i, h)
    style_header(ws, 5, 4)
    ref = {}
    r = 6
    for code, (tc, pc, gl) in MPFS.items():
        put(ws, r, 1, code)
        for j, v in enumerate((tc, pc, gl), 2):
            put(ws, r, j, v, F_INPUT, USD)
        ref[("mpfs", code)] = r
        r += 1
    put(ws, r, 1, "Full battery 95924 + 95923", F_BOLD)
    r924, r923 = ref[("mpfs", "95924")], ref[("mpfs", "95923")]
    for j, col in ((2, "B"), (3, "C"), (4, "D")):
        put(ws, r, j, f"={col}{r924}+{col}{r923}", F_BOLD, USD, FILL_KEY)
    ref["battery_row"] = r
    ws.cell(r, 4).comment = Comment("User-stated global 316.15 = 174.79 + 141.36 (independent-site "
                                    "Medicare global for the battery). Recomputed here from the rows above.", "build")
    ws.cell(r, 3).comment = Comment("User-stated 142.57 = 95.23 + 47.34. Used as the approximate "
                                    "physician fee: the physician group's own rate is not in hospital MRFs.", "build")
    r += 1
    put(ws, r, 1, "User-stated battery totals (check): global 316.15, technical 173.58, professional 142.57",
        F_NOTE)
    put(ws, r + 1, 1, "Check (should all be 0):", F_NOTE)
    put(ws, r + 1, 2, f"=ROUND(B{ref['battery_row']}-{MPFS_BATTERY_GIVEN['technical']},2)", F_NOTE, USD)
    put(ws, r + 1, 3, f"=ROUND(C{ref['battery_row']}-{MPFS_BATTERY_GIVEN['professional']},2)", F_NOTE, USD)
    put(ws, r + 1, 4, f"=ROUND(D{ref['battery_row']}-{MPFS_BATTERY_GIVEN['global']},2)", F_NOTE, USD)
    PHYS = f"Reference!$C${ref['battery_row']}"
    GLOBAL = f"Reference!$D${ref['battery_row']}"

    r += 4
    put(ws, r, 1, "CY2026 OPPS Addendum B (July 2026 update) — NATIONAL payment rates, "
                  "NOT adjusted for the Massachusetts wage index", F_BOLD)
    put(ws, r + 1, 1, f"Source: CMS, July 2026 OPPS Addendum B (updated July 21, 2026), file "
                      f"'2026 July Web Addendum B.07.14.26.xlsx'. {ADDB_URL}", F_NOTE)
    r += 2
    for i, h in enumerate(["Code", "Short descriptor", "Status indicator", "APC",
                           "National payment rate", "SI meaning"], 1):
        put(ws, r, i, h)
    style_header(ws, r, 6)
    si_text = {"S": "Paid separately; not discounted when multiple",
               "Q1": "STV-packaged: packaged (no separate payment) when on the same claim as an "
                     "S, T or V code; otherwise paid via its APC",
               "T": "Paid separately; multiple procedure reduction applies"}
    r += 1
    opps_first = r
    for code in CODES:
        a = addb[code]
        put(ws, r, 1, code)
        put(ws, r, 2, a["desc"])
        put(ws, r, 3, a["si"], F_INPUT)
        put(ws, r, 4, a["apc"], F_INPUT)
        put(ws, r, 5, a["rate"], F_INPUT, USD)
        put(ws, r, 6, si_text.get(a["si"], ""), wrap=True)
        ref[("opps", code)] = r
        r += 1
    opps_last = r - 1
    OPPS_CODES = f"Reference!$A${opps_first}:$A${opps_last}"
    OPPS_RATES = f"Reference!$E${opps_first}:$E${opps_last}"
    r += 1
    put(ws, r, 1, "OPPS labour-related share (assumption, for implied wage index only)", F_BOLD)
    put(ws, r, 5, 0.6, F_INPUT, "0.00")
    ws.cell(r, 5).comment = Comment("CMS applies the wage index to 60% of the OPPS rate "
                                    "(labour-related share). Used only in Medicare_check.", "build")
    LABOUR = f"Reference!$E${r}"
    widths(ws, {"A": 34, "B": 30, "C": 18, "D": 14, "E": 20, "F": 60})

    # ---------------------------------------------------------------- Rates_data
    ws = S["Rates_data"]
    cols = ["row_id", "hospital", "hosp", "code", "payer_name", "payer_family", "segment",
            "plan_name", "setting", "billing_class", "methodology", "rate_type", "gross_charge",
            "standard_charge_dollar", "standard_charge_percentage", "implied_dollar",
            "implied_flag", "pct_based", "median_allowed", "p10_allowed", "p90_allowed",
            "allowed_count", "headline_rate", "headline_basis", "included", "exclusion_reason",
            "medicare_comparator", "comparator_source", "multiple_vs_comparator", "review_flag",
            "opps_national", "multiple_vs_opps", "algorithm_text", "source_ref"]
    C = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
    for i, c in enumerate(cols, 1):
        put(ws, 1, i, c)
    style_header(ws, 1, len(cols))
    n = len(df)
    last = n + 1
    rng = lambda c: f"Rates_data!${C[c]}$2:${C[c]}${last}"  # noqa: E731
    # Medicare FFS row per (hospital, code) for the comparator
    mcr = {}
    for i, rw in df.iterrows():
        if rw.segment == "Medicare FFS" and rw.included:
            mcr.setdefault((rw.hospital, rw.code), i + 2)
    rowmap = {}
    for i, rw in df.iterrows():
        x = i + 2
        rowmap[(rw.hospital, rw.payer_name, rw.plan_name, rw.code, rw.setting)] = x
        vals = [i + 1, rw.hospital, SHORT[rw.hospital], rw.code, rw.payer_name, rw.family, rw.segment,
                rw.plan_name, rw.setting, rw.billing_class, rw.methodology, rw.rate_type,
                fnum(rw.gross_charge), fnum(rw.standard_charge_dollar),
                fnum(rw.standard_charge_percentage), fnum(rw.implied_dollar), rw.implied_flag,
                int(rw.pct_based), fnum(rw.median_amount), fnum(rw.percentile_10th),
                fnum(rw.percentile_90th), rw.allowed_count]
        for j, v in enumerate(vals, 1):
            put(ws, x, j, v, fmt=USD if cols[j - 1] in (
                "gross_charge", "standard_charge_dollar", "implied_dollar", "median_allowed",
                "p10_allowed", "p90_allowed") else None)
        P, S_, Y = f"{C['implied_dollar']}{x}", f"{C['median_allowed']}{x}", f"{C['included']}{x}"
        put(ws, x, cols.index("headline_rate") + 1, f'=IF({S_}<>"",{S_},IF({P}<>"",{P},""))', fmt=USD)
        put(ws, x, cols.index("headline_basis") + 1,
            f'=IF({S_}<>"","median allowed",IF({P}<>"","negotiated",""))')
        put(ws, x, cols.index("included") + 1, int(rw.included))
        put(ws, x, cols.index("exclusion_reason") + 1, rw.exclusion)
        # comparator: own hospital's Medicare FFS row for this code, else OPPS national
        if (rw.hospital, rw.code) in mcr:
            put(ws, x, cols.index("medicare_comparator") + 1,
                f"={C['implied_dollar']}{mcr[(rw.hospital, rw.code)]}", F_LINK, USD)
            put(ws, x, cols.index("comparator_source") + 1, "hospital's Medicare FFS row (APC)")
        else:
            put(ws, x, cols.index("medicare_comparator") + 1,
                f"=Reference!$E${ref[('opps', rw.code)]}", F_LINK, USD)
            put(ws, x, cols.index("comparator_source") + 1, "OPPS Addendum B national (no Medicare FFS row in file)")
        AA, AC = f"{C['medicare_comparator']}{x}", f"{C['multiple_vs_comparator']}{x}"
        put(ws, x, cols.index("multiple_vs_comparator") + 1,
            f'=IF(AND({Y}=1,{P}<>"",{AA}>0),{P}/{AA},"")', fmt=MULT)
        put(ws, x, cols.index("review_flag") + 1,
            f'=IF({AC}="","",IF(OR({AC}>20,{AC}<0.2),"REVIEW","ok"))')
        put(ws, x, cols.index("opps_national") + 1,
            f"=INDEX({OPPS_RATES},MATCH({C['code']}{x},{OPPS_CODES},0))", F_LINK, USD)
        AE = f"{C['opps_national']}{x}"
        put(ws, x, cols.index("multiple_vs_opps") + 1,
            f'=IF(AND({Y}=1,{P}<>"",{AE}>0),{P}/{AE},"")', fmt=MULT)
        put(ws, x, cols.index("algorithm_text") + 1, rw.standard_charge_algorithm[:300])
        put(ws, x, cols.index("source_ref") + 1, rw.source_ref)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = {
            "hospital": 30, "payer_name": 30, "plan_name": 36, "exclusion_reason": 40,
            "comparator_source": 34, "algorithm_text": 50, "source_ref": 50, "payer_family": 26,
            "segment": 22, "methodology": 24}.get(c, 13)

    def crit(**kw):
        """COUNTIFS-style criteria string over Rates_data."""
        return ",".join(f"{rng(k)},{v}" for k, v in kw.items())

    # ---------------------------------------------------------------- By_code_payer
    ws = S["By_code_payer"]
    title(ws, "Outpatient facility rates by hospital, code and payer (plans grouped within payer and segment)",
          "Included rows only (see QA). Headline = median allowed amount where the file gives one "
          "(what was actually paid), otherwise the median negotiated rate. All cells are formulas over Rates_data.")
    hdr = ["Hospital", "Code", "Payer (as in file)", "Payer family", "Segment", "Plan count",
           "Min implied $", "Median implied $", "Max implied $", "Share percentage-based",
           "Median allowed (median of plan medians)", "Lowest 10th pct allowed",
           "Highest 90th pct allowed", "Allowed counts (as in file)", "Headline rate", "Headline basis"]
    for i, h in enumerate(hdr, 1):
        put(ws, 4, i, h)
    style_header(ws, 4, len(hdr), 45)
    inc = df[df.included == 1]
    groups = inc.groupby(["hospital", "code", "segment", "payer_name"], sort=False)
    r = 5
    for (hosp, code, seg, payer), g in groups:
        put(ws, r, 1, hosp)
        put(ws, r, 2, code)
        put(ws, r, 3, payer)
        put(ws, r, 4, g.family.iloc[0])
        put(ws, r, 5, seg)
        k = crit(hospital=f"$A{r}", code=f"$B{r}", payer_name=f"$C{r}", segment=f"$E{r}", included=1)
        cond = (f"({rng('hospital')}=$A{r})*({rng('code')}=$B{r})*({rng('segment')}=$E{r})*"
                f"({rng('payer_name')}=$C{r})*({rng('included')}=1)")
        put(ws, r, 6, f"=COUNTIFS({k})")
        put(ws, r, 7, f"=_xlfn.MINIFS({rng('implied_dollar')},{k})", fmt=USD)
        ws.cell(r, 8).value = ArrayFormula(
            f"H{r}", f"=MEDIAN(IF({cond}*({rng('implied_dollar')}<>\"\"),{rng('implied_dollar')}))")
        ws.cell(r, 8).number_format, ws.cell(r, 8).font = USD, F_BASE
        put(ws, r, 9, f"=_xlfn.MAXIFS({rng('implied_dollar')},{k})", fmt=USD)
        put(ws, r, 10, f"=IF(F{r}=0,\"\",SUMIFS({rng('pct_based')},{k})/F{r})", fmt=PCT)
        nmed = f"COUNTIFS({k},{rng('median_allowed')},\"<>\")"
        ws.cell(r, 11).value = ArrayFormula(
            f"K{r}", f"=IF({nmed}=0,\"\",MEDIAN(IF({cond}*({rng('median_allowed')}<>\"\"),"
                     f"{rng('median_allowed')})))")
        ws.cell(r, 11).number_format, ws.cell(r, 11).font = USD, F_BASE
        put(ws, r, 12, f"=IF(COUNTIFS({k},{rng('p10_allowed')},\"<>\")=0,\"\","
                       f"_xlfn.MINIFS({rng('p10_allowed')},{k}))", fmt=USD)
        put(ws, r, 13, f"=IF(COUNTIFS({k},{rng('p90_allowed')},\"<>\")=0,\"\","
                       f"_xlfn.MAXIFS({rng('p90_allowed')},{k}))", fmt=USD)
        put(ws, r, 14, "; ".join(c for c in g.allowed_count if c) or "")
        put(ws, r, 15, f'=IF(K{r}<>"",K{r},H{r})', F_BOLD, USD)
        put(ws, r, 16, f'=IF(K{r}<>"","median allowed","median negotiated")')
        r += 1
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:P{r - 1}"
    widths(ws, {"A": 32, "B": 8, "C": 32, "D": 28, "E": 22, "F": 8, "G": 12, "H": 12, "I": 12,
                "J": 12, "K": 14, "L": 12, "M": 12, "N": 26, "O": 13, "P": 18})

    # ---------------------------------------------------------------- Battery
    ws = S["Battery"]
    title(ws, "Full autonomic battery: 95924 + 95923 facility rate, same payer/plan, plus approximate physician fee",
          "Physician fee = Medicare PC 142.57 (Reference). It is an APPROXIMATION: the physician "
          "group's negotiated rate is not published in hospital MRFs.")
    hdr = ["Hospital", "Segment", "Payer", "Plan", "Pricing basis",
           "95924 negotiated", "95923 negotiated", "Facility battery (negotiated)",
           "95924 median allowed", "95923 median allowed", "Facility battery (allowed)",
           "Headline facility battery", "Headline basis",
           "Approx. physician fee (Medicare PC, approximation)", "Approx. total battery cost", "Note"]
    for i, h in enumerate(hdr, 1):
        put(ws, 4, i, h)
    style_header(ws, 4, len(hdr), 45)
    r = 5
    battery_rows = []  # (row, hospital, segment)
    for hosp, short, _ in HOSPITALS:
        # Medicare FFS reference rows
        m924 = mcr.get((hosp, "95924"))
        m923 = mcr.get((hosp, "95923"))
        opps924, opps923 = f"Reference!$E${ref[('opps', '95924')]}", f"Reference!$E${ref[('opps', '95923')]}"
        variants = []
        if m924:
            variants.append(("Hospital Medicare FFS row, OPPS packaging applied (95923 = 0)",
                             f"=Rates_data!{C['implied_dollar']}{m924}", "=0",
                             "95923 is SI Q1: packaged with 95924 (SI S) on the same claim, no separate "
                             "payment. Faulkner's own 95923 row says it is only paid if highest rate."))
            if m923:
                variants.append(("Hospital Medicare FFS rows, as listed (no packaging)",
                                 f"=Rates_data!{C['implied_dollar']}{m924}",
                                 f"=Rates_data!{C['implied_dollar']}{m923}",
                                 "Sum of listed APC prices; overstates Medicare payment when billed together."))
            else:
                variants.append(("Hospital 95924 row + OPPS national 95923 (not listed in file)",
                                 f"=Rates_data!{C['implied_dollar']}{m924}", f"={opps923}",
                                 "File lists no Medicare row for 95923; national unadjusted Addendum B "
                                 "rate substituted and flagged."))
        if any(h == hosp for h, _ in mcr) or hosp in inc.hospital.values:
            variants.append(("OPPS national (Addendum B), packaging applied (95923 = 0)",
                             f"={opps924}", "=0", "National, unadjusted for MA wage index."))
            variants.append(("OPPS national (Addendum B), as listed", f"={opps924}", f"={opps923}",
                             "National, unadjusted for MA wage index; no packaging."))
        for basis, f924, f923, note in variants:
            put(ws, r, 1, hosp)
            put(ws, r, 2, "Medicare FFS")
            put(ws, r, 3, "Medicare" if "row" in basis else "Medicare (reference)")
            put(ws, r, 4, "")
            put(ws, r, 5, basis, wrap=True)
            put(ws, r, 6, f924, F_LINK, USD)
            put(ws, r, 7, f923, F_LINK, USD)
            put(ws, r, 8, f"=F{r}+G{r}", fmt=USD)
            if basis.startswith("Hospital"):  # file rows may carry median allowed amounts
                ma = C["median_allowed"]
                put(ws, r, 9, f'=IF(Rates_data!{ma}{m924}="","",Rates_data!{ma}{m924})', F_LINK, USD)
                if "packaging applied" in basis:
                    put(ws, r, 10, "=0", F_LINK, USD)
                elif m923:
                    put(ws, r, 10, f'=IF(Rates_data!{ma}{m923}="","",Rates_data!{ma}{m923})', F_LINK, USD)
                put(ws, r, 11, f'=IF(AND(I{r}<>"",J{r}<>""),I{r}+J{r},"")', fmt=USD)
            put(ws, r, 12, f'=IF(K{r}<>"",K{r},H{r})', F_BOLD, USD)
            put(ws, r, 13, f'=IF(K{r}<>"","median allowed","APC price")')
            put(ws, r, 14, f"={PHYS}", F_LINK, USD)
            put(ws, r, 15, f"=L{r}+N{r}", F_BOLD, USD)
            put(ws, r, 16, note, F_NOTE, wrap=True)
            battery_rows.append((r, hosp, "Medicare FFS", basis))
            r += 1
        # payer/plan pairs with both codes
        pairs = {}
        for rw in inc[(inc.hospital == hosp) & (inc.code.isin(BATTERY))
                      & (inc.segment != "Medicare FFS")].itertuples():
            pairs.setdefault((rw.segment, rw.payer_name, rw.plan_name), {}).setdefault(rw.code, []).append(rw)
        seg_order = {"Commercial": 0, "Medicare Advantage": 1, "MassHealth / dual managed care": 2, "TRICARE": 3}
        for (seg, payer, plan), d in sorted(pairs.items(), key=lambda kv: (seg_order.get(kv[0][0], 9), kv[0][1], kv[0][2])):
            if not all(c in d for c in BATTERY):
                continue
            if any(len(d[c]) > 1 for c in BATTERY):
                raise SystemExit(f"ambiguous pairing (several settings) for {hosp} {payer} {plan}")
            x924 = rowmap[(hosp, payer, plan, "95924", d["95924"][0].setting)]
            x923 = rowmap[(hosp, payer, plan, "95923", d["95923"][0].setting)]
            put(ws, r, 1, hosp)
            put(ws, r, 2, seg)
            put(ws, r, 3, payer)
            put(ws, r, 4, plan)
            put(ws, r, 5, "Hospital file, same payer/plan")
            put(ws, r, 6, f"=Rates_data!{C['implied_dollar']}{x924}", F_LINK, USD)
            put(ws, r, 7, f"=Rates_data!{C['implied_dollar']}{x923}", F_LINK, USD)
            put(ws, r, 8, f"=F{r}+G{r}", fmt=USD)
            put(ws, r, 9, f'=IF(Rates_data!{C["median_allowed"]}{x924}="","",Rates_data!{C["median_allowed"]}{x924})', F_LINK, USD)
            put(ws, r, 10, f'=IF(Rates_data!{C["median_allowed"]}{x923}="","",Rates_data!{C["median_allowed"]}{x923})', F_LINK, USD)
            put(ws, r, 11, f'=IF(AND(I{r}<>"",J{r}<>""),I{r}+J{r},"")', fmt=USD)
            put(ws, r, 12, f'=IF(K{r}<>"",K{r},H{r})', F_BOLD, USD)
            put(ws, r, 13, f'=IF(K{r}<>"","median allowed","negotiated")')
            put(ws, r, 14, f"={PHYS}", F_LINK, USD)
            put(ws, r, 15, f"=L{r}+N{r}", F_BOLD, USD)
            note = ""
            if seg == "Medicare Advantage" and d["95923"][0].standard_charge_algorithm:
                note = "APC-priced MA plan; 95923 likely packaged in practice (Q1)."
            put(ws, r, 16, note, F_NOTE, wrap=True)
            battery_rows.append((r, hosp, seg, "pair"))
            r += 1
    bat_first, bat_last = 5, r - 1
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:P{bat_last}"
    # Summary block
    r += 2
    put(ws, r, 1, "Commercial summary by hospital (medians across commercial payer/plan pairs)", F_BOLD)
    r += 1
    sh = ["Hospital", "Commercial pairs", "Median headline facility battery",
          "Median negotiated facility battery", "Median approx. total battery cost", "Min headline", "Max headline"]
    for i, h in enumerate(sh, 1):
        put(ws, r, i, h)
    style_header(ws, r, len(sh), 40)
    r += 1
    bsum = {}
    A_, B_ = f"$A${bat_first}:$A${bat_last}", f"$B${bat_first}:$B${bat_last}"
    for hosp, short, _ in HOSPITALS:
        cond = f"({A_}=$A{r})*({B_}=\"Commercial\")"
        put(ws, r, 1, hosp)
        put(ws, r, 2, f'=COUNTIFS({A_},$A{r},{B_},"Commercial")')
        for col, src in ((3, "L"), (4, "H"), (5, "O")):
            ws.cell(r, col).value = ArrayFormula(
                f"{get_column_letter(col)}{r}",
                f'=IF($B{r}=0,"none listed",MEDIAN(IF({cond},${src}${bat_first}:${src}${bat_last})))')
            ws.cell(r, col).number_format, ws.cell(r, col).font = USD, F_BOLD if col == 3 else F_BASE
        put(ws, r, 6, f'=IF($B{r}=0,"",_xlfn.MINIFS($L${bat_first}:$L${bat_last},{A_},$A{r},{B_},"Commercial"))', fmt=USD)
        put(ws, r, 7, f'=IF($B{r}=0,"",_xlfn.MAXIFS($L${bat_first}:$L${bat_last},{A_},$A{r},{B_},"Commercial"))', fmt=USD)
        bsum[hosp] = r
        r += 1
    widths(ws, {"A": 32, "B": 18, "C": 30, "D": 36, "E": 34, "F": 12, "G": 12, "H": 13, "I": 12,
                "J": 12, "K": 13, "L": 14, "M": 15, "N": 16, "O": 14, "P": 48})

    # ---------------------------------------------------------------- Multiples
    ws = S["Multiples"]
    title(ws, "Multiples of Medicare",
          "A: each included facility rate / OPPS national (Addendum B, unadjusted for MA wage index) and / "
          "the hospital's own Medicare FFS row. B: approximate total battery cost / 316.15 (independent-site MPFS global).")
    put(ws, 4, 1, "A. Facility rate multiples", F_BOLD)
    hdr = ["Hospital", "Code", "Segment", "Payer", "Plan", "Implied facility rate",
           "OPPS national rate", "Multiple of OPPS national", "Hospital Medicare FFS row (same code)",
           "Multiple of hospital Medicare row", "Headline rate", "Headline basis"]
    for i, h in enumerate(hdr, 1):
        put(ws, 5, i, h)
    style_header(ws, 5, len(hdr), 40)
    r = 6
    for i, rw in df.iterrows():
        if not rw.included:
            continue
        x = i + 2
        put(ws, r, 1, rw.hospital)
        put(ws, r, 2, rw.code)
        put(ws, r, 3, rw.segment)
        put(ws, r, 4, rw.payer_name)
        put(ws, r, 5, rw.plan_name)
        put(ws, r, 6, f"=Rates_data!{C['implied_dollar']}{x}", F_LINK, USD)
        put(ws, r, 7, f"=Rates_data!{C['opps_national']}{x}", F_LINK, USD)
        put(ws, r, 8, f"=IF(G{r}>0,F{r}/G{r},\"\")", fmt=MULT)
        if (rw.hospital, rw.code) in mcr:
            put(ws, r, 9, f"=Rates_data!{C['implied_dollar']}{mcr[(rw.hospital, rw.code)]}", F_LINK, USD)
            put(ws, r, 10, f"=IF(I{r}>0,F{r}/I{r},\"\")", fmt=MULT)
        else:
            put(ws, r, 9, "not listed")
            put(ws, r, 10, "")
        put(ws, r, 11, f"=Rates_data!{C['headline_rate']}{x}", F_LINK, USD)
        put(ws, r, 12, f"=Rates_data!{C['headline_basis']}{x}", F_LINK)
        r += 1
    r += 2
    put(ws, r, 1, "B. Battery: approximate total cost (hospital facility + 142.57 physician approximation) "
                  "/ 316.15 independent-site Medicare global", F_BOLD)
    r += 1
    hdr = ["Hospital", "Segment", "Payer", "Plan", "Pricing basis", "Approx. total battery cost",
           "MPFS global (independent site)", "Site-of-service multiple"]
    for i, h in enumerate(hdr, 1):
        put(ws, r, i, h)
    style_header(ws, r, len(hdr), 40)
    r += 1
    mult_bat = {}
    for br, hosp, seg, basis in battery_rows:
        for j, col in enumerate("ABCDE", 1):
            put(ws, r, j, f'=IF(Battery!{col}{br}="","",Battery!{col}{br})', F_LINK)
        put(ws, r, 6, f"=Battery!O{br}", F_LINK, USD)
        put(ws, r, 7, f"={GLOBAL}", F_LINK, USD)
        put(ws, r, 8, f"=F{r}/G{r}", F_BOLD, MULT)
        mult_bat[br] = r
        r += 1
    ws.freeze_panes = "C6"
    widths(ws, {"A": 32, "B": 18, "C": 22, "D": 32, "E": 40, "F": 14, "G": 14, "H": 12,
                "I": 16, "J": 14, "K": 12, "L": 16})

    # ---------------------------------------------------------------- Focus_payers
    ws = S["Focus_payers"]
    title(ws, "Focus payers — all included rows",
          "MassHealth managed care: only Commonwealth Care Alliance (One Care / SCO dual plans) appears; "
          "no MassHealth ACO or WellSense rows exist for these codes in any file.")
    hdr = ["Focus group", "Hospital", "Code", "Segment", "Payer", "Plan", "Setting", "Methodology",
           "Negotiated / implied $", "Percentage of charges", "Median allowed", "10th pct allowed",
           "90th pct allowed", "Allowed count", "Headline rate", "Headline basis",
           "Medicare comparator", "Multiple of comparator"]
    for i, h in enumerate(hdr, 1):
        put(ws, 4, i, h)
    style_header(ws, 4, len(hdr), 40)
    r = 5
    for label, colkey, val, seg in FOCUS:
        field = "family" if colkey == "F" else "segment"
        sub = inc[(inc[field] == val) & ((inc.segment == seg) if seg else True)]
        if sub.empty:
            put(ws, r, 1, label, F_BOLD)
            put(ws, r, 2, "not listed at any hospital for these codes", F_NOTE)
            r += 1
            continue
        for i, rw in sub.iterrows():
            x = i + 2
            put(ws, r, 1, label, F_BOLD)
            for j, v in enumerate([rw.hospital, rw.code, rw.segment, rw.payer_name, rw.plan_name,
                                   rw.setting, rw.methodology], 2):
                put(ws, r, j, v)
            for j, c in enumerate(["implied_dollar", "standard_charge_percentage", "median_allowed",
                                   "p10_allowed", "p90_allowed", "allowed_count", "headline_rate",
                                   "headline_basis", "medicare_comparator", "multiple_vs_comparator"], 9):
                put(ws, r, j, f'=IF(Rates_data!{C[c]}{x}="","",Rates_data!{C[c]}{x})', F_LINK,
                    USD if c in ("implied_dollar", "median_allowed", "p10_allowed", "p90_allowed",
                                 "headline_rate", "medicare_comparator")
                    else MULT if c == "multiple_vs_comparator" else None)
            r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:R{r - 1}"
    widths(ws, {"A": 30, "B": 30, "C": 8, "D": 20, "E": 30, "F": 34, "G": 11, "H": 24})

    # ---------------------------------------------------------------- Coverage
    ws = S["Coverage"]
    title(ws, "Coverage: does the hospital file list rows for 95923 / 95924?",
          "Commercial rows for the seven commercial focus payers; any rows for MassHealth managed care, "
          "Medicare Advantage and Medicare FFS. 'not listed' = no row in the hospital's MRF.")
    put(ws, 4, 1, "Focus payer")
    put(ws, 4, 2, "Rows counted")
    c = 3
    cov_cols = []
    for hosp, short, _ in HOSPITALS:
        for code in BATTERY[::-1]:
            put(ws, 4, c, f"{short} {code}")
            cov_cols.append((c, hosp, code))
            c += 1
    style_header(ws, 4, c - 1)
    r = 5
    rows = FOCUS + [("Medicare fee-for-service", "G", "Medicare FFS", None)]
    for label, colkey, val, seg in rows:
        put(ws, r, 1, label, F_BOLD)
        put(ws, r, 2, f"{seg} rows" if seg else f"{val} rows")
        for c, hosp, code in cov_cols:
            k = f"{rng('hospital')},\"{hosp}\",{rng('code')},\"{code}\",{rng('included')},1," + (
                f"{rng('payer_family')},\"{val}\"" if colkey == "F" else f"{rng('segment')},\"{val}\"")
            if seg:
                k += f",{rng('segment')},\"{seg}\""
            put(ws, r, c, f'=IF(COUNTIFS({k})=0,"not listed",COUNTIFS({k})&" row(s)")')
        r += 1
    r += 1
    notes = [
        "MGH's file lists NO Medicare fee-for-service row for 95923, although Medicare pays for 95923 "
        "(Q1: separately payable when not billed with an S/T/V code; Faulkner lists an APC price of $153.08). "
        "MGH's de-identified minimum for 95923 ($145.42) is below every listed payer rate, consistent with an "
        "unlisted low (Medicare-type) payer.",
        "Brigham and Women's Hospital and Wentworth-Douglass Hospital list no rows for any of the five codes, "
        "for any payer.",
        "Blue Cross Blue Shield of MA and Harvard Pilgrim appear commercially only at BIDMC. In MGB files BCBS "
        "appears only on a Medicare (Advantage) plan; Harvard Pilgrim does not appear for these codes.",
        "Tufts Health Plan appears only on Medicare Advantage ('Medicare replacement') plans.",
    ]
    for t in notes:
        put(ws, r, 1, t, F_NOTE)
        r += 1
    widths(ws, {"A": 38, "B": 32, **{get_column_letter(i): 13 for i in range(3, 13)}})

    # ---------------------------------------------------------------- Reconciliation
    ws = S["Reconciliation"]
    title(ws, "Reconciliation: Medicare FFS rows in hospital files vs CY2024 Medicare claims patient counts",
          "Patient counts supplied by the user (blue). MRF 'count' = number of allowed amounts in the file's "
          "look-back period; it counts claims/lines, not patients, and the period is not necessarily CY2024.")
    hdr = ["Hospital (clinicians)", "Code", "CY2024 Medicare patients (user)", "Medicare FFS rows in MRF",
           "MRF allowed-amount count (Medicare FFS)", "MRF Medicare FFS rate", "Medicare Advantage rows in MRF",
           "Assessment"]
    for i, h in enumerate(hdr, 1):
        put(ws, 4, i, h)
    style_header(ws, 4, len(hdr), 45)
    r = 5
    assess = {
        ("MGH", "95924"): "Consistent in scale: 84 allowed amounts vs 94 patients (different period and unit).",
        ("MGH", "95923"): "Gap: 95 Medicare patients but no Medicare row. Likely because 95923 (Q1) is packaged "
                          "into 95924 on the same claim, leaving no separate allowed amount to publish.",
        ("BWFH", "95924"): "Consistent: 121 allowed amounts vs 114 patients.",
        ("BWFH", "95923"): "Row listed but count only '1 through 10' vs 114 patients: consistent with 95923 "
                           "being packaged on most claims.",
        ("BIDMC", "95924"): "Gap: 99 Medicare patients but BIDMC's file lists no Medicare rows for any "
                            "target code (only BCBS, Harvard Pilgrim, Unicare).",
        ("BIDMC", "95923"): "Gap: as for 95924.",
    }
    for hosp, short, label in HOSPITALS:
        if not label:
            continue
        for code in BATTERY:
            put(ws, r, 1, label)
            put(ws, r, 2, code)
            put(ws, r, 3, PATIENTS_2024[(short, code)], F_INPUT)
            k = f"{rng('hospital')},\"{hosp}\",{rng('code')},\"{code}\",{rng('segment')},\"Medicare FFS\""
            put(ws, r, 4, f"=COUNTIFS({k})")
            x = mcr.get((hosp, code))
            if x:
                put(ws, r, 5, f"=Rates_data!{C['allowed_count']}{x}", F_LINK)
                put(ws, r, 6, f"=Rates_data!{C['implied_dollar']}{x}", F_LINK, USD)
            else:
                put(ws, r, 5, "not listed")
                put(ws, r, 6, "not listed")
            put(ws, r, 7, f"=COUNTIFS({rng('hospital')},\"{hosp}\",{rng('code')},\"{code}\","
                          f"{rng('segment')},\"Medicare Advantage\")")
            put(ws, r, 8, assess[(short, code)], wrap=True)
            r += 1
    widths(ws, {"A": 26, "B": 8, "C": 14, "D": 12, "E": 16, "F": 14, "G": 14, "H": 80})

    # ---------------------------------------------------------------- Medicare_check
    ws = S["Medicare_check"]
    title(ws, "Medicare FFS rows in hospital files (primary comparator) vs OPPS Addendum B (national)",
          "Addendum B rates are national and unadjusted for the Massachusetts wage index. Implied wage "
          "index = (file price / national - (1 - labour share)) / labour share.")
    hdr = ["Hospital", "Code", "MRF Medicare APC price", "APC price in algorithm text",
           "MRF median allowed", "MRF allowed count", "Addendum B SI", "Addendum B APC",
           "Addendum B national rate", "Difference $", "Ratio file / national", "Implied wage index",
           "Comment"]
    for i, h in enumerate(hdr, 1):
        put(ws, 4, i, h)
    style_header(ws, 4, len(hdr), 45)
    r = 5
    comments = {
        "95921": "Implied wage index inconsistent with the 95923/95924 rows; file APC price not reproducible "
                 "from the CY2026 APC with one wage index.",
        "95923": "Close to a Boston-area wage adjustment of APC 5734. Median allowed is 2x the APC price "
                 "(possibly 2 units or mixed claims).",
        "95924": "File price ($369.8) is 1.68x the national APC 5722 rate; would need wage index ~2.1. "
                 "Closer to APC 5723 ($381.24 national). The hospital's APC assignment or rate year may "
                 "differ from CY2026 Addendum B. Median allowed confirms ~$370 actually paid.",
    }
    for (hosp, code), x in sorted(mcr.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        rw = df.iloc[x - 2]
        m = re.search(r"APC Price \(\$([\d,.]+)\)", rw.standard_charge_algorithm)
        put(ws, r, 1, hosp)
        put(ws, r, 2, code)
        put(ws, r, 3, f"=Rates_data!{C['implied_dollar']}{x}", F_LINK, USD)
        put(ws, r, 4, float(m.group(1).replace(",", "")) if m else None, fmt=USD)
        put(ws, r, 5, f'=IF(Rates_data!{C["median_allowed"]}{x}="","",Rates_data!{C["median_allowed"]}{x})', F_LINK, USD)
        put(ws, r, 6, f"=Rates_data!{C['allowed_count']}{x}", F_LINK)
        put(ws, r, 7, f"=Reference!$C${ref[('opps', code)]}", F_LINK)
        put(ws, r, 8, f"=Reference!$D${ref[('opps', code)]}", F_LINK)
        put(ws, r, 9, f"=Reference!$E${ref[('opps', code)]}", F_LINK, USD)
        put(ws, r, 10, f"=C{r}-I{r}", fmt=USD)
        put(ws, r, 11, f"=C{r}/I{r}", fmt=MULT)
        put(ws, r, 12, f"=(K{r}-(1-{LABOUR}))/{LABOUR}", fmt="0.00")
        put(ws, r, 13, comments.get(code, ""), F_NOTE, wrap=True)
        r += 1
    for hosp, short, _ in HOSPITALS:
        if hosp in inc.hospital.values and not any(h == hosp for h, _ in mcr):
            put(ws, r, 1, hosp)
            put(ws, r, 2, "all")
            put(ws, r, 13, "No Medicare FFS rows in this hospital's file; Addendum B national used as comparator.",
                F_NOTE, wrap=True)
            r += 1
    widths(ws, {"A": 32, "B": 8, "C": 13, "D": 13, "E": 13, "F": 12, "G": 9, "H": 9, "I": 13,
                "J": 12, "K": 11, "L": 11, "M": 70})

    # ---------------------------------------------------------------- QA
    ws = S["QA"]
    title(ws, "QA: exclusions, missing codes, review flags")
    r = 3
    put(ws, r, 1, "1. Rows excluded from payer statistics", F_BOLD)
    r += 1
    for i, h in enumerate(["Row id", "Hospital", "Code", "Payer", "Plan", "Reason"], 1):
        put(ws, r, i, h)
    style_header(ws, r, 6)
    r += 1
    for i, rw in df[df.included == 0].iterrows():
        for j, v in enumerate([i + 1, rw.hospital, rw.code, rw.payer_name or "(none)",
                               rw.plan_name or "(none)", rw.exclusion], 1):
            put(ws, r, j, v)
        r += 1
    put(ws, r, 1, "Also not in Battery: rows whose payer/plan lacks either 95923 or 95924, and 93660 / 95921 / "
                  "95922 (not part of the battery). They remain in By_code_payer and Multiples.", F_NOTE)
    r += 2
    put(ws, r, 1, "2. Included rows per hospital and code ('missing' = no included row)", F_BOLD)
    r += 1
    put(ws, r, 1, "Hospital")
    for j, code in enumerate(CODES, 2):
        put(ws, r, j, code)
    style_header(ws, r, len(CODES) + 1)
    r += 1
    for hosp, short, _ in HOSPITALS:
        put(ws, r, 1, hosp)
        for j, code in enumerate(CODES, 2):
            k = f"{rng('hospital')},$A{r},{rng('code')},\"{code}\",{rng('included')},1"
            put(ws, r, j, f'=IF(COUNTIFS({k})=0,"missing",COUNTIFS({k}))')
        r += 1
    r += 1
    put(ws, r, 1, "3. Rates above 20x or below 0.2x their Medicare comparator", F_BOLD)
    r += 1
    put(ws, r, 1, "Rows flagged REVIEW:")
    put(ws, r, 2, f'=COUNTIFS({rng("review_flag")},"REVIEW")', F_BOLD)
    put(ws, r, 3, "Rows checked:")
    put(ws, r, 4, f'=COUNTIFS({rng("review_flag")},"ok")+COUNTIFS({rng("review_flag")},"REVIEW")')
    r += 1
    put(ws, r, 1, "Highest multiple of comparator:")
    put(ws, r, 2, f"=MAX({rng('multiple_vs_comparator')})", fmt=MULT)
    put(ws, r, 3, "Lowest:")
    put(ws, r, 4, f"=MIN({rng('multiple_vs_comparator')})", fmt=MULT)
    r += 1
    put(ws, r, 1, "Comparator = the hospital's own Medicare FFS row for the code where listed, otherwise OPPS "
                  "Addendum B national. Filter Rates_data!review_flag for the rows.", F_NOTE)
    r += 2
    put(ws, r, 1, "4. Data caveats found during QA", F_BOLD)
    r += 1
    for t in [
        "MGB allowed-amount medians appear pooled above plan level: UnitedHealthcare 'New Business Discount' and "
        "'PPO/POS' share identical medians (761.83 / 889.73) despite different negotiated rates.",
        "Some median allowed amounts exceed the negotiated rate (e.g. MGH Aetna 95924: allowed 1,444.72 vs "
        "negotiated 801.30); may reflect multiple units per claim or a different plan mix. Headline figures "
        "use median allowed where present, as specified.",
        "MGB files have no billing_class column; rows treated as hospital facility (all MGB plan names are "
        "'HB' = hospital billing).",
        "BIDMC payer rows carry no gross charge and sit on a separate item from the gross/cash row; no rate "
        "was derived across items.",
        "No percentage-only rates exist in these files, so implied_flag is never TRUE.",
        "Medicare FFS APC prices in MGB files do not reconcile to CY2026 Addendum B with a single wage index "
        "(see Medicare_check).",
    ]:
        put(ws, r, 1, "• " + t, F_NOTE)
        r += 1
    widths(ws, {"A": 34, "B": 32, "C": 12, "D": 34, "E": 36, "F": 60})

    # ---------------------------------------------------------------- Notes
    ws = ws_notes
    title(ws, "Autonomic function testing — hospital price transparency summary")
    lines = [
        ("Scope", "CPT 95921, 95922, 95923, 95924 (autonomic function tests) and 93660 (tilt table, context). "
                  "Hospitals: MGH, Brigham and Women's, BW Faulkner, Wentworth-Douglass, BIDMC."),
        ("Source", "Hospital machine-readable files (CMS HPT, schema v3.0.0) found via each system's cms-hpt.txt; "
                   "normalised in hpt_autonomic_rates.csv (Rates_data tab)."),
        ("IMPORTANT", "These are HOSPITAL FACILITY rates (hospital outpatient department). They bound "
                      "independent-site rates from above. They are NOT AutoLabs' expected rates."),
        ("OPPS Addendum B", "CY2026 OPPS Addendum B (July 2026 update) rates are NATIONAL and UNADJUSTED for the "
                            "Massachusetts wage index."),
        ("Physician fee", "The 142.57 physician component is the Medicare MPFS professional component "
                          "(95924 PC 95.23 + 95923 PC 47.34), used as an APPROXIMATION: physician-group rates "
                          "are not in hospital files."),
        ("Headline figures", "Where a row carries median / 10th / 90th percentile allowed amounts and a count, "
                             "the median allowed amount (what was actually paid) is the headline; otherwise "
                             "the negotiated rate."),
        ("Medicare packaging", "95923 has OPPS status indicator Q1: when billed on the same claim as 95924 (SI S) "
                               "it is packaged and not paid separately. Battery tab shows Medicare both with "
                               "packaging applied and as listed."),
        ("Comparator", "Medicare facility comparator = the hospital's own Medicare FFS (APC-priced) row where "
                       "listed; OPPS Addendum B national otherwise (always for BIDMC). Cross-check in "
                       "Medicare_check."),
        ("Colour key", "Blue text = input value; green text = link to another sheet; black = formula; "
                       "yellow fill = key derived total."),
        ("Tabs", "By_code_payer · Battery · Multiples · Focus_payers · Coverage · Reconciliation · "
                 "Medicare_check · QA · Reference · Rates_data"),
    ]
    for i, (k, v) in enumerate(lines, 3):
        put(ws, i, 1, k, F_BOLD)
        put(ws, i, 2, v, fill=FILL_KEY if k == "IMPORTANT" else None, wrap=True)
        ws.row_dimensions[i].height = 30
    widths(ws, {"A": 20, "B": 120})

    for s in wb.worksheets:
        s.sheet_view.zoomScale = 90
    wb.save(args.out)
    print(f"wrote {args.out}: {len(df)} data rows, {len(battery_rows)} battery rows")


if __name__ == "__main__":
    main()
