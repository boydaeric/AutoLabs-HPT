#!/usr/bin/env python3
"""Extract text from every comment's attachments (OCR where needed) and merge it
with the comment text.

Reads  data/<docket>/<comment>.json  and  data/<docket>/attachments/*
Writes data/<docket>/_text/<comment>.json   {merged_text, comment_text, attachments:[...], problems:[...]}

PDFs: pdftotext; any page with almost no text layer is rasterised and OCR'd with
tesseract. DOCX: python-docx (paragraphs + tables). Images: tesseract.
Where one attachment exists in several formats (pdf + docx) only one is used so the
text isn't double counted (docx preferred). Idempotent: existing _text files are kept
unless --force.
"""
import argparse
import json
import re
import subprocess
import tempfile
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

DATA = Path("data")
MIN_PAGE_CHARS = 40       # below this a PDF page is treated as scanned
MIN_DOC_CHARS = 30        # below this an attachment counts as "no usable text"
OCR_DPI = "200"
FORMAT_PREF = {"docx": 0, "doc": 1, "pdf": 2}


def run(cmd, timeout):
    return subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)


def ocr_image(path):
    r = run(["tesseract", str(path), "stdout", "-l", "eng"], 300)
    return r.stdout


STOPWORDS = {"the", "of", "and", "to", "in", "a", "is", "that", "for", "on", "with", "as", "are", "this", "be", "we", "by", "it"}


def garbled(text):
    """True when a page has a text layer but it is mojibake (broken font encoding) rather than English prose."""
    w = re.findall(r"[A-Za-z]+", text.lower())
    if len(w) < 30 and text.count("\u00ff") < 10:
        return False
    return (sum(x in STOPWORDS for x in w) / max(1, len(w)) < 0.07) or text.count("\u00ff") / max(1, len(text)) > 0.01


def pdf_text(path):
    info = run(["pdfinfo", str(path)], 60)
    if info.returncode != 0:
        raise RuntimeError("pdfinfo failed: " + info.stderr.strip()[:200])
    pages = int(next((l.split()[1] for l in info.stdout.splitlines() if l.startswith("Pages:")), 0))
    if "Encrypted:      yes" in info.stdout and "copy:no" in info.stdout.replace(" ", ""):
        pass  # pdftotext will still try; failure is caught below
    out, ocr_pages = [], 0
    # pdftotext emits \f between pages
    r = run(["pdftotext", "-layout", str(path), "-"], 300)
    if r.returncode != 0:
        raise RuntimeError("pdftotext failed: " + r.stderr.strip()[:200])
    page_txt = r.stdout.split("\f")
    if page_txt and not page_txt[-1].strip():
        page_txt = page_txt[:-1]
    page_txt += [""] * (pages - len(page_txt))
    for i, t in enumerate(page_txt[:pages], 1):
        if len(t.strip()) >= MIN_PAGE_CHARS and not garbled(t):
            out.append(t)
            continue
        with tempfile.TemporaryDirectory() as td:
            base = Path(td) / "p"
            rr = run(["pdftoppm", "-r", OCR_DPI, "-f", str(i), "-l", str(i), "-png", str(path), str(base)], 300)
            imgs = sorted(Path(td).glob("p*.png"))
            if rr.returncode != 0 or not imgs:
                out.append(t)
                continue
            out.append(ocr_image(imgs[0]))
            ocr_pages += 1
    return "\n".join(out), {"pages": pages, "ocr_pages": ocr_pages}


def docx_text(path):
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            seen = []
            for c in row.cells:
                if c.text not in seen:
                    seen.append(c.text)
            parts.append(" | ".join(seen))
    return "\n".join(parts), {}


def extract_file(path):
    ext = path.suffix.lower().lstrip(".")
    if ext == "pdf":
        return pdf_text(path)
    if ext == "docx":
        return docx_text(path)
    if ext in ("png", "jpg", "jpeg", "tif", "tiff"):
        return ocr_image(path), {"ocr_pages": 1}
    if ext in ("doc", "rtf", "odt", "txt"):
        r = run(["pandoc", str(path), "-t", "plain"], 120)
        if r.returncode != 0:
            raise RuntimeError("pandoc failed: " + r.stderr.strip()[:200])
        return r.stdout, {}
    raise RuntimeError(f"unsupported format {ext}")


def process(args):
    rec_path, force = args
    rec = json.loads(rec_path.read_text())
    out_dir = rec_path.parent / "_text"
    out_path = out_dir / rec_path.name
    if out_path.exists() and not force:
        return rec["id"], "cached"
    # one file per attachment slot (attachment_N), preferring docx over pdf
    slots = {}
    for att in rec["attachments"]:
        if not att.get("url"):
            continue
        slot = att["url"].rsplit("/", 1)[-1].rsplit(".", 1)[0]
        cur = slots.get(slot)
        if cur is None or FORMAT_PREF.get(att["format"], 9) < FORMAT_PREF.get(cur["format"], 9):
            slots[slot] = att
    results, problems, texts = [], [], []
    for slot, att in sorted(slots.items()):
        path = Path(att["local_path"]) if att.get("local_path") else None
        entry = {"file": path.name if path else None, "format": att["format"]}
        try:
            if not path or not path.exists():
                raise RuntimeError("attachment file missing")
            text, meta = extract_file(path)
            text = text.replace("\x00", "")
            entry.update(meta, chars=len(text.strip()), status="ok")
            if len(text.strip()) < MIN_DOC_CHARS:
                entry["status"] = "no_text"
                problems.append(f"{entry['file']}: no usable text extracted")
            else:
                texts.append(text.strip())
        except Exception as e:  # noqa: BLE001 - record every failure, keep going
            entry.update(status="error", error=str(e)[:300], chars=0)
            problems.append(f"{entry['file']}: {e}"[:300])
        results.append(entry)
    comment_text = (rec.get("comment") or "").strip()
    if not comment_text and not texts:
        problems.append("no comment text and no attachment text")
    merged = "\n\n".join(([comment_text] if comment_text else []) + texts)
    out_dir.mkdir(exist_ok=True)
    out_path.write_text(json.dumps({
        "id": rec["id"], "comment_text": comment_text, "attachments": results,
        "problems": problems, "merged_text": merged}, ensure_ascii=False))
    return rec["id"], "problems" if problems else "ok"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--workers", type=int, default=4)
    a = ap.parse_args()
    paths = sorted(p for p in DATA.glob("*/CMS-*.json"))
    counts = {}
    with ProcessPoolExecutor(a.workers) as ex:
        for n, (cid, status) in enumerate(ex.map(process, [(p, a.force) for p in paths], chunksize=4), 1):
            counts[status] = counts.get(status, 0) + 1
            if n % 100 == 0:
                print(f"{n}/{len(paths)} {counts}", flush=True)
    print("done", len(paths), counts)


if __name__ == "__main__":
    main()
